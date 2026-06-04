import csv
import logging
from datetime import datetime, date

from django.http import HttpResponse
from django_filters.rest_framework import DjangoFilterBackend
from rest_framework import viewsets, filters, status
from rest_framework.decorators import action
from rest_framework.permissions import AllowAny
from rest_framework.response import Response
from rest_framework.views import APIView
from rest_framework_simplejwt.tokens import RefreshToken
from rest_framework_simplejwt.views import TokenObtainPairView
from rest_framework.authentication import SessionAuthentication
from rest_framework_simplejwt.authentication import JWTAuthentication

from drf_yasg.utils import swagger_auto_schema
from drf_yasg import openapi

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

from .filters import (
    AutorFilter, CategoriaFilter, EditorialFilter,
    UsuarioFilter, LibroFilter, PrestamoFilter,
    DevolucionFilter, MultaFilter,
)
from .models import Autor, Categoria, Editorial, Usuario, Libro, Prestamo, Devolucion, Multa
from .pagination import StandardPagination
from .serializers import (
    AutorSerializer, CategoriaSerializer, EditorialSerializer,
    UsuarioSerializer, LibroSerializer, PrestamoSerializer,
    DevolucionSerializer, MultaSerializer,
)
from .utils import success_response, error_response, created_response

logger = logging.getLogger('biblioteca')


def log_operation(user, action, model, object_id, details):
    username = user.username if user and user.is_authenticated else "anonymous"
    logger.info(f"[{action}] User={username} | Model={model} | ID={object_id} | Details={details}")


# ==========================================
# AUTH: Login con logging
# ==========================================
class CustomTokenObtainPairView(TokenObtainPairView):
    def post(self, request, *args, **kwargs):
        response = super().post(request, *args, **kwargs)
        if response.status_code == 200:
            username = request.data.get('username', '')
            logger.info(f"[LOGIN] user={username}")
            return Response({
                "success": True,
                "message": "Autenticación exitosa",
                "data": response.data,
            }, status=status.HTTP_200_OK)
        return response


class LogoutView(APIView):
    def post(self, request):
        try:
            refresh_token = request.data.get("refresh")
            if not refresh_token:
                return error_response("Token de refresco requerido")
            token = RefreshToken(refresh_token)
            token.blacklist()
            logger.info(f"[LOGOUT] user={request.user.username if request.user.is_authenticated else 'anonymous'}")
            return success_response("Sesión cerrada correctamente")
        except Exception as e:
            return error_response(str(e))


# ==========================================
# MIXIN BASE para todos los ViewSets
# ==========================================
class BaseViewSet(viewsets.ModelViewSet):
    pagination_class = StandardPagination
    filter_backends = [DjangoFilterBackend, filters.OrderingFilter, filters.SearchFilter]
    authentication_classes = [JWTAuthentication, SessionAuthentication]

    # Para el listado general, ocultamos los desactivados
    def get_queryset(self):
        return super().get_queryset().filter(activo=True)

    def get_permissions(self):
        return [AllowAny()]

    # 👇 SOLUCIÓN CLAVE: Permite ver el registro individual aunque esté inactivo (Soft Delete)
    def retrieve(self, request, *args, **kwargs):
        lookup_url_kwarg = self.lookup_url_kwarg or self.lookup_field
        filter_kwargs = {self.lookup_field: self.kwargs[lookup_url_kwarg]}
        
        try:
            # Buscamos en todo el universo de objetos (.all()), ignorando el filtro de activo=True
            instance = self.queryset.model.objects.get(**filter_kwargs)
            serializer = self.get_serializer(instance)
            return Response(serializer.data, status=status.HTTP_200_OK)
            
        except self.queryset.model.DoesNotExist:
            return Response({
                'success': False,
                'status': 404,
                'version': 'v1',
                'message': 'El registro solicitado no existe en el sistema.'
            }, status=status.HTTP_404_NOT_FOUND)

    @swagger_auto_schema(
        method='get',
        operation_description="Exporta todos los registros activos de la tabla actual a un documento estructurado de Excel.",
        responses={
            200: openapi.Response(
                description="Archivo Excel descargado exitosamente (.xlsx)",
                schema=openapi.Schema(type=openapi.TYPE_FILE)
            ),
            404: "No hay registros para exportar"
        }
    )
    @action(detail=False, methods=['get'], url_path='export/excel')
    def export_excel(self, request):
        if not OPENPYXL_AVAILABLE:
            return error_response("La librería 'openpyxl' no está instalada en el entorno virtual.")

        queryset = self.filter_queryset(self.get_queryset())
        
        if not queryset.exists():
            return error_response("No se encontraron registros disponibles para realizar la exportación.", status_code=status.HTTP_404_NOT_FOUND)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = getattr(self, 'basename', 'Datos').capitalize()

        font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        fill_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        alignment_center = Alignment(horizontal="center", vertical="center")

        model = queryset.model
        fields = [f for f in model._meta.get_fields() if f.concrete and not f.many_to_many]

        headers = [f.name.replace('_', ' ').upper() for f in fields]
        ws.append(headers)

        for cell in ws[1]:
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = alignment_center

        for obj in queryset:
            row_data = []
            for field in fields:
                val = getattr(obj, field.name)
                
                if isinstance(val, datetime):
                    val = val.replace(tzinfo=None)
                elif isinstance(val, date):
                    val = val.strftime('%Y-%m-%d')
                elif field.is_relation and val is not None:
                    val = str(val)
                elif isinstance(val, bool):
                    val = "SÍ" if val else "NO"
                
                row_data.append(val if val is not None else "")
            ws.append(row_data)

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        filename = f"{ws.title.lower()}_export.xlsx"
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        wb.save(response)
        
        return response

    # 👇 MODIFICADO: Retorna los datos completos con activo=False tras la desactivación
    def destroy(self, request, pk=None, *args, **kwargs):
        if not request.user or not request.user.is_authenticated or not request.user.is_superuser:
            return Response({
                'success': False,
                'status': 403,
                'version': 'v1',
                'message': 'Acceso denegado. No cuenta con los privilegios de Administrador requeridos para eliminar registros.'
            }, status=status.HTTP_403_FORBIDDEN)

        try:
            # Buscamos de la base de datos completa para evitar el filtro del queryset original
            instance = self.queryset.model.objects.get(pk=pk)
            instance.activo = False
            instance.save()

            log_operation(
                user=request.user,
                action='SOFT_DELETE',
                model=self.queryset.model.__name__,
                object_id=instance.id,
                details={'activo': False}
            )

            logger.info(f"SOFT_DELETE: {self.queryset.model.__name__} ID={instance.id}")
            
            # Serializamos la instancia con sus datos completos actualizados (activo=False)
            serializer = self.get_serializer(instance)
            return Response(serializer.data, status=status.HTTP_200_OK)

        except self.queryset.model.DoesNotExist:
            return Response({
                'success': False,
                'status': 404,
                'version': 'v1',
                'message': 'El registro especificado no existe o ya ha sido removido del sistema.'
            }, status=status.HTTP_404_NOT_FOUND)
        except Exception as e:
            return Response({
                'success': False,
                'status': 400,
                'version': 'v1',
                'message': f'Error al procesar la solicitud: {str(e)}'
            }, status=status.HTTP_400_BAD_REQUEST)

    def list(self, request, *args, **kwargs):
        queryset = self.filter_queryset(self.get_queryset())
        page = self.paginate_queryset(queryset)
        if page is not None:
            serializer = self.get_serializer(page, many=True)
            return self.get_paginated_response(serializer.data)
        serializer = self.get_serializer(queryset, many=True)
        return success_response(data=serializer.data)

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        if serializer.is_valid():
            instance = serializer.save()
            username = request.user.username if request.user.is_authenticated else "anonymous"
            logger.info(f"[CREACION] Tabla={self.basename} | ID={instance.id} | user={username}")
            return created_response(data=serializer.data)
        return error_response("Error de validación en los datos enviados.", errors=serializer.errors)

    def update(self, request, *args, **kwargs):
        partial = kwargs.pop('partial', False)
        
        # Obtenemos el objeto directamente saltando el filtro activo=True por si requiere modificarse
        lookup_url_kwarg = self.lookup_url_kwarg or self.lookup_field
        filter_kwargs = {self.lookup_field: self.kwargs[lookup_url_kwarg]}
        instance = self.queryset.model.objects.get(**filter_kwargs)
        
        serializer = self.get_serializer(instance, data=request.data, partial=partial)
        if serializer.is_valid():
            serializer.save()
            username = request.user.username if request.user.is_authenticated else "anonymous"
            logger.info(f"[MODIFICACION] Tabla={self.basename} | ID={instance.id} | user={username}")
            return success_response("El registro ha sido actualizado correctamente.", data=serializer.data)
        return error_response("Error de validación en los campos modificados.", errors=serializer.errors)


# ==========================================
# VISTAS DE LOS MODELOS
# ==========================================

class AutorViewSet(BaseViewSet):
    queryset = Autor.objects.all()
    serializer_class = AutorSerializer
    filterset_class = AutorFilter
    search_fields = ['nombre', 'biografia']
    ordering_fields = ['nombre', 'id']


class CategoriaViewSet(BaseViewSet):
    queryset = Categoria.objects.all()
    serializer_class = CategoriaSerializer
    filterset_class = CategoriaFilter
    search_fields = ['nombre', 'descripcion']
    ordering_fields = ['nombre', 'id']


class EditorialViewSet(BaseViewSet):
    queryset = Editorial.objects.all()
    serializer_class = EditorialSerializer
    filterset_class = EditorialFilter
    search_fields = ['nombre', 'direccion']
    ordering_fields = ['nombre', 'id']


class UsuarioViewSet(BaseViewSet):
    queryset = Usuario.objects.all()
    serializer_class = UsuarioSerializer
    filterset_class = UsuarioFilter
    search_fields = ['nombre', 'email', 'num_tarjeta']
    ordering_fields = ['nombre', 'fecha_registro', 'id']


class LibroViewSet(BaseViewSet):
    queryset = Libro.objects.all()
    serializer_class = LibroSerializer
    filterset_class = LibroFilter
    search_fields = ['titulo', 'isbn']
    ordering_fields = ['titulo', 'anio_publicacion', 'id']


class PrestamoViewSet(BaseViewSet):
    queryset = Prestamo.objects.all()
    serializer_class = PrestamoSerializer
    filterset_class = PrestamoFilter
    search_fields = ['usuario__nombre', 'libro__titulo']
    ordering_fields = ['fecha_prestamo', 'fecha_devolucion_esperada', 'id']


class DevolucionViewSet(BaseViewSet):
    queryset = Devolucion.objects.all()
    serializer_class = DevolucionSerializer
    filterset_class = DevolucionFilter
    search_fields = ['prestamo__usuario__nombre', 'prestamo__libro__titulo']
    ordering_fields = ['fecha_devolucion_real', 'id']


class MultaViewSet(BaseViewSet):
    queryset = Multa.objects.all()
    serializer_class = MultaSerializer
    filterset_class = MultaFilter
    search_fields = ['usuario__nombre']
    ordering_fields = ['monto', 'estado', 'id']